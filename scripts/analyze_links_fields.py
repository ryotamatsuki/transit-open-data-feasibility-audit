#!/usr/bin/env python3
"""Inspect the acquired LINKS CSV schemas and report row-level completeness."""
from __future__ import annotations

import csv
import io
import re
import zipfile
from pathlib import Path

from common import ANALYSIS, RAW, links_dataset_files, normalize_name, parse_date, write_csv

FIELDS = {
    "operator_name": ("事業者名", "事業者", "会社名"),
    "operator_number": ("事業者番号", "事業者コード", "事業者ID"),
    "year": ("年度", "年次", "対象年度"),
    "permit_type": ("許可区分", "事業区分", "許可種別", "事業種別"),
    "prefecture": ("都道府県", "都道府県名", "管轄運輸"),
    "address": ("住所", "所在地", "事業者基本情報_住所"),
    "vehicle_count": ("車両数", "事業用自動車数"),
    "extended_actual_vehicle_count": ("延実在車両数",),
    "extended_working_vehicle_count": ("延実働車両数",),
    "actual_vehicle_km": ("実車キロ", "実車キロ数"),
    "total_vehicle_km": ("走行キロ", "走行距離"),
    "passenger_count": ("輸送人員", "輸送人数", "旅客数"),
    "revenue": ("営業収入", "運送収入", "運輸収入"),
    "sales_revenue": ("営業収入", "営業収益"),
    "transport_revenue": ("運送収入", "運輸収入"),
    "operating_expenses": ("営業費用", "営業経費"),
    "driver_count": ("運転者数", "運転者人数", "バス運転者数"),
    "employee_count": ("従業員数", "社員数", "職員数"),
}


def csv_records_from(path: Path):
    if path.suffix.lower() == ".csv":
        with path.open("rb") as f:
            raw = f.read()
        text = raw.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text, newline=""))
        yield path.name, reader.fieldnames or [], list(reader)
    elif path.suffix.lower() == ".zip":
        with zipfile.ZipFile(path) as zf:
            for name in zf.namelist():
                if name.lower().endswith(".csv"):
                    with zf.open(name) as f:
                        text = io.TextIOWrapper(f, encoding="utf-8-sig", errors="replace", newline="")
                        reader = csv.DictReader(text)
                        yield name, reader.fieldnames or [], list(reader)
    elif path.suffix.lower() in (".xlsx", ".xls"):
        if path.suffix.lower() == ".xls":
            import xlrd
            workbook = xlrd.open_workbook(path, on_demand=True)
            try:
                for sheet in workbook.sheets():
                    columns = [str(x).strip() for x in sheet.row_values(0)] if sheet.nrows else []
                    rows = []
                    for row_idx in range(1, sheet.nrows):
                        values = sheet.row_values(row_idx)
                        rows.append({columns[i]: str(values[i]) if i < len(values) and values[i] is not None else "" for i in range(len(columns)) if columns[i]})
                    yield f"{path.name}::{sheet.name}", columns, rows
            finally:
                workbook.release_resources()
        else:
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in workbook.worksheets:
                    iterator = sheet.iter_rows(values_only=True)
                    header = next(iterator, ())
                    columns = [str(x).strip() if x is not None else "" for x in header]
                    rows = []
                    for values in iterator:
                        padded = list(values) + [None] * max(0, len(columns) - len(values))
                        rows.append({columns[i]: "" if padded[i] is None else str(padded[i]) for i in range(len(columns)) if columns[i]})
                    yield f"{path.name}::{sheet.title}", columns, rows
            finally:
                workbook.close()


def classify_file(name: str) -> str:
    s = name.lower()
    if "概況" in name or "jigyougaiyou" in s or "busreport" in s:
        return "business_overview"
    if "輸送実績" in name or "yusoujisseki" in s or "transport" in s:
        return "transport_performance"
    if "specification" in s or "仕様書" in name:
        return "specification"
    return "other"


def has_value(value: object) -> bool:
    return value is not None and str(value).strip() not in ("", "-", "－", "NA", "N/A", "null", "None")


def infer_year(value: str) -> str:
    m = re.search(r"(?:19|20)\d{2}", value or "")
    return m.group(0) if m else ""


def main() -> None:
    files = links_dataset_files()
    outputs: list[dict[str, object]] = []
    raw_rows: list[dict[str, object]] = []
    operator_by_id: dict[str, dict[str, object]] = {}
    if not files:
        for field, _terms in FIELDS.items():
            outputs.append({"dataset": "not_acquired", "source_file": "", "year": "", "field": field, "columns": "", "rows_total": "", "rows_nonblank": "", "unique_operator_ids_nonblank": "", "status": "not_evaluable_source_not_retrieved", "grain_note": "No Project LINKS CSV bytes were acquired; blank counts are unknown, not zero."})
        write_csv(ANALYSIS / "field_completeness.csv", outputs)
        write_csv(ANALYSIS / "links_operator_metrics.csv", [], ["links_operator_id", "operator_name", "operator_number", "prefecture_address", "years", "has_transport_performance", "has_business_overview", "has_driver_count", "source_rows"])
        write_csv(ANALYSIS / "links_raw_rows.csv", [], ["source_file", "source_row", "dataset", "operator_name", "operator_number", "year", "permit_type", "prefecture_address", "passenger_count", "vehicle_count", "revenue", "driver_count"])
        print("LINKS data unavailable: no raw CSV/ZIP files found; field counts remain unknown.")
        return

    for file in files:
        for member, columns, rows in csv_records_from(file):
            dataset = classify_file(member)
            field_columns = {
                field: [c for c in columns
                        if any(t in c for t in terms)
                        and not (field == "vehicle_count" and any(mark in c for mark in ("延実在", "延実働")))]
                for field, terms in FIELDS.items()
            }
            name_col = next(iter(field_columns["operator_name"]), "")
            number_col = next(iter(field_columns["operator_number"]), "")
            year_col = next(iter(field_columns["year"]), "")
            permit_col = next(iter(field_columns["permit_type"]), "")
            loc_col = next(iter(field_columns["address"]), "") or next(iter(field_columns["prefecture"]), "")
            for row_num, row in enumerate(rows, 1):
                raw_name = str(row.get(name_col, "")) if name_col else ""
                number = str(row.get(number_col, "")) if number_col else ""
                year = infer_year(str(row.get(year_col, ""))) if year_col else ""
                permit = str(row.get(permit_col, "")) if permit_col else ""
                location = str(row.get(loc_col, "")) if loc_col else ""
                # Retain every row and its source identity; never drop repeated
                # company/year/type records during grain assessment.
                if raw_name or number:
                    identity_source = number.strip() or (normalize_name(raw_name) + "|" + normalize_name(location))
                    oid = "L-" + __import__("hashlib").sha256(identity_source.encode("utf-8")).hexdigest()[:16]
                    op = operator_by_id.setdefault(oid, {"links_operator_id": oid, "operator_name": raw_name, "operator_number": number, "prefecture_address": location, "years": set(), "datasets": set(), "source_rows": 0, "has_driver_count": False, "has_transport_performance": False, "has_business_overview": False})
                    if year:
                        op["years"].add(year)
                    op["datasets"].add(dataset)
                    op["source_rows"] += 1
                    op["has_driver_count"] |= any(has_value(row.get(c)) for c in field_columns["driver_count"])
                    op["has_transport_performance"] |= dataset == "transport_performance"
                    op["has_business_overview"] |= dataset == "business_overview"
                else:
                    oid = ""
                raw_rows.append({"source_file": member, "source_row": row_num, "dataset": dataset, "operator_name": raw_name, "operator_number": number, "year": year, "permit_type": permit, "prefecture_address": location, "passenger_count": "|".join(str(row.get(c, "")) for c in field_columns["passenger_count"]), "vehicle_count": "|".join(str(row.get(c, "")) for c in field_columns["vehicle_count"]), "revenue": "|".join(str(row.get(c, "")) for c in field_columns["revenue"]), "driver_count": "|".join(str(row.get(c, "")) for c in field_columns["driver_count"]), "links_operator_id": oid})
            years = sorted({infer_year(str(r.get(year_col, ""))) for r in rows}) if year_col else [""]
            years = [x for x in years if x] or [""]
            for year in years:
                year_rows = [r for r in rows if not year or (year_col and infer_year(str(r.get(year_col, ""))) == year)]
                for field, cols in field_columns.items():
                    nonblank = [r for r in year_rows if any(has_value(r.get(c)) for c in cols)]
                    ids = {"L-" + __import__("hashlib").sha256(((str(r.get(number_col, "")).strip() if number_col else "") or (normalize_name(str(r.get(name_col, ""))) + "|" + normalize_name(str(r.get(loc_col, ""))))).encode("utf-8")).hexdigest()[:16] for r in nonblank if name_col and has_value(r.get(name_col))}
                    outputs.append({"dataset": dataset, "source_file": member, "year": year, "field": field, "columns": "|".join(cols), "rows_total": len(year_rows), "rows_nonblank": len(nonblank), "unique_operator_ids_nonblank": len(ids), "status": "measured_from_csv" if cols else "column_not_present", "grain_note": "Source records retained row by row; repeated operator/year/permit rows are not dropped."})

    operator_rows = []
    for op in operator_by_id.values():
        op = dict(op)
        op["years"] = "|".join(sorted(op["years"]))
        op["datasets"] = "|".join(sorted(op["datasets"]))
        operator_rows.append(op)
    write_csv(ANALYSIS / "field_completeness.csv", outputs)
    write_csv(ANALYSIS / "links_operator_metrics.csv", operator_rows)
    write_csv(ANALYSIS / "links_raw_rows.csv", raw_rows)
    print(f"LINKS CSV files: {len(files)}; rows: {len(raw_rows)}; provisional company keys: {len(operator_rows)}")


if __name__ == "__main__":
    main()
